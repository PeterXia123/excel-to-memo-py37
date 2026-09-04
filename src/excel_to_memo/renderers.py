from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Tuple, Union

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.shared import Inches, RGBColor

from .transform import MemoReport, build_comments_text

ACCENT = RGBColor(138, 69, 69)
TEXT = RGBColor(25, 25, 25)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EDEDED"
TABLE_FILL = "F3F3F3"
BORDER = "B7B7B7"
REQUIREMENT_NOTES = {
    "Monitoring Frequency meets minimum requirements": [
        "(High Risk = Quarterly; Medium Risk = Semi-Annually; Low Risk = Annually; or earlier as agreed-upon)"
    ],
    "Nature and type of monitoring is aligned with type of model and its output characteristics": [
        "(i.e., monitoring is consistent with agreed-upon methodology)"
    ],
}


def render_markdown(report: MemoReport, output_path: Union[str, Path]) -> Path:
    path = Path(output_path)
    lines: List[str] = []
    lines.append(f"# {report.title}")
    lines.append("")
    lines.append("## Monitoring Report Summary")
    lines.append("")
    lines.append("| Field | Value |")
    lines.append("| --- | --- |")
    for key, value in report.monitoring_summary.items():
        lines.append(f"| {key} | {value} |")
    lines.append("")
    lines.append("## MVA Review Summary")
    lines.append("")
    lines.append("| Field | Value |")
    lines.append("| --- | --- |")
    for key, value in report.review_summary.items():
        lines.append(f"| {key} | {value} |")
    lines.append("")
    lines.append("## MVA Assessment")
    lines.append("")
    lines.append("### Model Monitoring Requirements")
    lines.append("")
    for item in report.requirements:
        lines.append(f"- [x] {item}")
    lines.append("")
    lines.append("### Comments")
    lines.append("")
    comments_text = build_comments_text(report)
    if comments_text:
        lines.extend(comments_text.splitlines())
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def render_docx(report: MemoReport, output_path: Union[str, Path]) -> Path:
    path = Path(output_path)
    document = Document()
    configure_document(document)

    add_title_block(document, report.title)

    add_monitoring_summary_block(document, report.monitoring_summary)
    add_review_summary_box(document, report.review_summary)

    add_section_heading(document, "MVA ASSESSMENT:")
    requirements = document.add_paragraph()
    requirements.paragraph_format.space_before = Pt(2)
    requirements.paragraph_format.space_after = Pt(6)
    label = requirements.add_run("Model Monitoring Requirements ")
    label.bold = True
    label.underline = True
    set_run_font(label, name="Arial", size=11.5, color=TEXT, bold=True)
    requirements.add_run("(select all that apply)")
    set_run_font(requirements.runs[-1], name="Arial", size=11, color=MUTED, italic=True)

    for item in report.requirements:
        add_requirement_item(document, item, REQUIREMENT_NOTES.get(item, []))

    add_comments_box(document, report)

    document.save(path)
    return path


def export_normalized_rows(report: MemoReport, output_path: Union[str, Path]) -> Path:
    path = Path(output_path)
    frame = pd.DataFrame(report.normalized_rows)
    frame.to_csv(path, index=False)
    return path


def add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, name="Arial", size=13, color=ACCENT, bold=True)
    run.underline = True


def add_monitoring_summary_block(document: Document, values: Dict[str, str]) -> None:
    add_section_heading(document, "MONITORING REPORT SUMMARY:")
    for key, value in values.items():
        add_tabbed_kv_paragraph(document, key, value, tab_stop=Inches(3.25), bold_value=True)


def add_review_summary_box(document: Document, values: Dict[str, str]) -> None:
    outer = document.add_table(rows=1, cols=1)
    outer.autofit = False
    set_table_indent(outer, 0)
    cell = outer.cell(0, 0)
    cell.width = Inches(6.25)
    set_cell_borders(cell, BORDER)
    set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
    if cell.paragraphs and not cell.paragraphs[0].text.strip():
        p = cell.paragraphs[0]._element
        p.getparent().remove(p)

    heading = cell.add_paragraph()
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(10)
    run = heading.add_run("MVA REVIEW SUMMARY:")
    set_run_font(run, name="Arial", size=13, color=ACCENT, bold=True)
    run.underline = True

    for key, value in values.items():
        italic_label = key == "Assessment Comments"
        add_tabbed_kv_paragraph(
            cell,
            key,
            value,
            tab_stop=Inches(3.45),
            bold_value=True,
            italic_label=italic_label,
        )


def add_key_value_table(
    document: Document,
    values: Dict[str, str],
    widths: Tuple[Inches, Inches],
    shaded: bool,
    compact: bool,
) -> None:
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    set_table_indent(table, 0)
    for key, value in values.items():
        row = table.add_row().cells
        row[0].width = widths[0]
        row[1].width = widths[1]
        write_cell_text(row[0], key, bold=True)
        write_cell_text(row[1], value)
        set_cell_borders(row[0], BORDER)
        set_cell_borders(row[1], BORDER)
        set_cell_margins(row[0], top=70, bottom=70, start=110, end=110)
        set_cell_margins(row[1], top=70, bottom=70, start=110, end=110)
        if shaded:
            shade_cell(row[0], LIGHT_FILL)
            shade_cell(row[1], LIGHT_FILL)
    if compact:
        apply_table_compact_style(table)
    else:
        apply_table_standard_style(table)


def add_tabbed_kv_paragraph(
    container,
    label: str,
    value: str,
    tab_stop,
    bold_value: bool = False,
    italic_label: bool = False,
) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.tab_stops.add_tab_stop(tab_stop)

    label_run = paragraph.add_run(f"{label}:")
    set_run_font(label_run, name="Arial", size=11.5, color=TEXT, bold=True, italic=italic_label)

    tab_run = paragraph.add_run("\t")
    set_run_font(tab_run, name="Arial", size=11.5, color=TEXT)

    value_run = paragraph.add_run(value)
    set_run_font(value_run, name="Arial", size=11.5, color=TEXT, bold=bold_value)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = TEXT
    normal_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.line_spacing = 1.08


def add_title_block(document: Document, title_text: str) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run(title_text)
    set_run_font(title_run, name="Arial", size=17, color=TEXT, bold=True)
    add_bottom_border(title, color="555555", size=10)


def add_comments_box(document: Document, report: MemoReport) -> None:
    table = document.add_table(rows=2, cols=1)
    table.autofit = False
    set_table_indent(table, 0)

    header_cell = table.cell(0, 0)
    body_cell = table.cell(1, 0)
    header_cell.width = Inches(6.2)
    body_cell.width = Inches(6.2)

    write_cell_text(header_cell, "Comments:", bold=True)
    shade_cell(header_cell, TABLE_FILL)
    set_cell_borders(header_cell, BORDER)
    set_cell_borders(body_cell, BORDER)
    set_cell_margins(header_cell, top=70, bottom=70, start=110, end=110)
    set_cell_margins(body_cell, top=90, bottom=90, start=130, end=130)

    for paragraph_text in report.intro_paragraphs:
        paragraph = body_cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(paragraph_text)
        set_run_font(run, name="Arial", size=10.5, color=TEXT)

    for bullet in report.bullets:
        paragraph = body_cell.add_paragraph()
        paragraph.style = document.styles["List Bullet"]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(0)
        apply_bold_label_bullet(paragraph, bullet)

    for paragraph_text in report.closing_paragraphs:
        paragraph = body_cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(paragraph_text)
        set_run_font(run, name="Arial", size=10.5, color=TEXT)

    # Remove the default empty paragraph that comes with new table cells.
    if body_cell.paragraphs and not body_cell.paragraphs[0].text.strip():
        p = body_cell.paragraphs[0]._element
        p.getparent().remove(p)


def apply_bold_label_bullet(paragraph, bullet_text: str) -> None:
    if ":" in bullet_text:
        label, rest = bullet_text.split(":", 1)
        label_run = paragraph.add_run(f"{label.strip()}:")
        set_run_font(label_run, name="Arial", size=10.5, color=TEXT, bold=True)
        spacer = paragraph.add_run(" ")
        set_run_font(spacer, name="Arial", size=10.5, color=TEXT)
        add_rich_body_runs(paragraph, rest.strip())
    else:
        add_rich_body_runs(paragraph, bullet_text)


def add_rich_body_runs(paragraph, body_text: str) -> None:
    sentence_pattern = re.compile(
        r"^The (?P<parameter>.+?) parameter(?: for (?P<portfolio>.+?))? "
        r"is (?P<direction>underestimated|overestimated) "
        r"(?:with materiality of (?P<materiality>.+?) )?"
        r"(?:and )?with an ECL impact of (?P<impact>.+?)\.$",
        re.IGNORECASE,
    )
    sentences = re.split(r"(?<=[.?!])\s+", body_text.strip())
    first = True
    for sentence in sentences:
        if not sentence:
            continue
        if not first:
            spacer = paragraph.add_run(" ")
            set_run_font(spacer, name="Arial", size=10.5, color=TEXT)
        first = False

        match = sentence_pattern.match(sentence.strip())
        if not match:
            run = paragraph.add_run(sentence.strip())
            set_run_font(run, name="Arial", size=10.5, color=TEXT)
            continue

        parameter = match.group("parameter")
        portfolio = match.group("portfolio")
        direction = match.group("direction")
        materiality = match.group("materiality")
        impact = match.group("impact")

        prefix = paragraph.add_run("The ")
        set_run_font(prefix, name="Arial", size=10.5, color=TEXT)

        param_run = paragraph.add_run(parameter)
        set_run_font(param_run, name="Arial", size=10.5, color=TEXT, bold=True)

        middle = " parameter"
        if portfolio:
            middle += f" for {portfolio}"
        middle += " is "
        middle_run = paragraph.add_run(middle)
        set_run_font(middle_run, name="Arial", size=10.5, color=TEXT)

        direction_run = paragraph.add_run(direction)
        set_run_font(direction_run, name="Arial", size=10.5, color=TEXT, bold=True)

        if materiality:
            materiality_bridge = paragraph.add_run(" with materiality of ")
            set_run_font(materiality_bridge, name="Arial", size=10.5, color=TEXT)

            materiality_run = paragraph.add_run(materiality)
            set_run_font(materiality_run, name="Arial", size=10.5, color=TEXT, bold=True)

            bridge_run = paragraph.add_run(" and an ECL impact of ")
        else:
            bridge_run = paragraph.add_run(" with an ECL impact of ")
        set_run_font(bridge_run, name="Arial", size=10.5, color=TEXT)

        impact_run = paragraph.add_run(impact)
        set_run_font(impact_run, name="Arial", size=10.5, color=TEXT, bold=True)

        period_run = paragraph.add_run(".")
        set_run_font(period_run, name="Arial", size=10.5, color=TEXT)


def add_requirement_item(document: Document, item: str, notes: List[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Inches(0.20)
    paragraph.paragraph_format.first_line_indent = Inches(0)

    checkbox = paragraph.add_run("☒ ")
    set_run_font(checkbox, name="Arial", size=10.5, color=TEXT)

    run = paragraph.add_run(item)
    set_run_font(run, name="Arial", size=10.5, color=TEXT, bold=True, italic=True)

    for note in notes:
        note_paragraph = document.add_paragraph()
        note_paragraph.paragraph_format.space_before = Pt(0)
        note_paragraph.paragraph_format.space_after = Pt(6)
        note_paragraph.paragraph_format.left_indent = Inches(0.42)
        note_run = note_paragraph.add_run(note)
        set_run_font(note_run, name="Arial", size=9.5, color=MUTED, italic=True)


def write_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, name="Arial", size=10.5, color=TEXT, bold=bold)


def set_run_font(run, name: str, size: float, color: RGBColor, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right"):
        element = tc_borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_bottom_border(paragraph, color: str, size: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)


def set_table_indent(table, indent_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def apply_table_standard_style(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")


def apply_table_compact_style(table) -> None:
    apply_table_standard_style(table)
